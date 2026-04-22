from odoo import models, fields
import openpyxl
import base64
import xlwt

from io import BytesIO
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class HrPayslipRun(models.Model):
    _inherit = 'hr.payslip.run'

    def print_xlsx_report2_kuec(self):
        """
                Generates two Excel worksheets for bank transfers:
                1. BT (Bank Transfer) - for employees whose bank matches the company bank.
                2. LBT (Local Bank Transfer) - for employees with a different bank.

                Both sheets include structured data fields like IBAN, SWIFT, bank address, transfer purpose, etc.
                """
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = 'BT'

        # Set column widths
        column_widths = [30, 10, 30, 30, 30, 30, 30, 30, 30, 30, 30, 30, 30, 30, 30, 30, 30, 30, 30,
                         30, 30, 30, 30]
        for i, width in enumerate(column_widths, start=1):
            worksheet.column_dimensions[get_column_letter(i)].width = width

        # Headers
        headers = [
            "Product Code", "Debit Account No.", "Beneficiary Nick Name", "Beneficiary Account No.",
            "Beneficiary Name", "Beneficiary Address Line 1", "Beneficiary Address Line 2",
            "Beneficiary Bank", "Sub_Bank_Counter", "Beneficiary Bank Location",
            "Beneficiary Bank Address", "SWIFT Address/BIC", "Clearing Code",
            "Payment Date (DD-MM-YYYY)", "Payment Currency", "Payment Amount",
            "Charge Type", "Purpose Code", "Purpose Of Payment",
            "Customer Reference", "Intermediate Bank Swift Code/BIC", "FX Contract Reference"
        ]

        side_fill = PatternFill(start_color="0155b8", end_color="0155b8", fill_type="solid")
        side_font = Font(color="FFFFFF", bold=True)
        side_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        worksheet.merge_cells('A1:A15')
        cell1 = worksheet.cell(row=1, column=1,
                               value="1. Own Account Transfer \n \n 2.Transfer Within Bank")
        cell1.fill = side_fill
        cell1.font = side_font
        cell1.alignment = side_alignment
        # worksheet.merge_cells('B1:B15')
        header_fill = PatternFill(start_color="ccdfe0", end_color="ccdfe0", fill_type="solid")
        header_font = Font(bold=True)
        header_alignment = Alignment(horizontal="center", vertical="center")

        for col_num, header in enumerate(headers, start=2):
            cell = worksheet.cell(row=1, column=col_num, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        # Sample data
        data = [
            ["BT", "Company FAB Account", "SS", "Employee FAB IBAN", "Employee complete name", "1",
             "2", "3", "X", "3", "3",
             "3", "11", "12", "13", "14", "1", "16", "17", "18", "19", "20"]
        ]

        # for row_num, row_data in enumerate(data, start=2):
        #     for col_num, cell_value in enumerate(row_data, start=2):
        #         worksheet.cell(row=row_num, column=col_num, value=cell_value)

        row = 2
        col = 2

        account = self.env.company.account_no or ""
        for slip in self.slip_ids:
            if slip.employee_id.bank_name_id and slip.employee_id.bank_name_id.id == self.env.company.bank_id.id:
                worksheet.cell(row=row, column=col, value="BT")
                col += 1
                worksheet.cell(row=row, column=col, value=account)
                col += 1

                worksheet.cell(row=row, column=col, value="")
                col += 1
                worksheet.cell(row=row, column=col, value=slip.employee_id.iban_num or "")
                col += 1
                worksheet.cell(row=row, column=col, value=slip.employee_id.display_name)
                col += 1
                worksheet.cell(row=row, column=col, value=slip.employee_id.addresline1 or "")
                col += 1
                worksheet.cell(row=row, column=col, value=slip.employee_id.addresline2 or "")
                col += 1

                worksheet.cell(row=row, column=col, value=slip.employee_id.bank_name_id.name or "")
                col += 1

                worksheet.cell(row=row, column=col, value="AE")
                col += 1
                worksheet.cell(row=row, column=col, value="UAE")
                col += 1
                worksheet.cell(row=row, column=col, value="UAE")
                col += 1

                worksheet.cell(row=row, column=col, value=slip.employee_id.swift_address_bic or "")
                col += 1
                worksheet.cell(row=row, column=col, value=self.date_end.strftime('%d/%m/%Y'))
                col += 1
                worksheet.cell(row=row, column=col, value=self.env.company.currency_id.display_name)
                col += 1

                worksheet.cell(row=row, column=col, value=slip.net_wage)
                col += 1

                worksheet.cell(row=row, column=col, value="OUR")
                col += 1
                worksheet.cell(row=row, column=col, value="")
                col += 1

                worksheet.cell(row=row, column=col,
                               value=self.date_end.strftime("%B %Y") + " Salary")
                col += 1

                col = 2
                row += 1

        worksheet_lbt = workbook.create_sheet(title='LBT')

        # Set column widths for LBT worksheet
        for i, width in enumerate(column_widths, start=1):
            worksheet_lbt.column_dimensions[get_column_letter(i)].width = width

        # Headers for LBT worksheet
        for col_num, header in enumerate(headers, start=2):
            cell = worksheet_lbt.cell(row=1, column=col_num, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        # Sample data for LBT worksheet
        # for row_num, row_data in enumerate(data, start=2):
        #     for col_num, cell_value in enumerate(row_data, start=2):
        #         worksheet_lbt.cell(row=row_num, column=col_num, value=cell_value)

        worksheet_lbt.merge_cells('A1:A15')
        cell22 = worksheet_lbt.cell(row=1, column=1,
                                    value="1. Own Account Transfer \n \n 2.Transfer Within Bank")
        cell22.fill = side_fill
        cell22.font = side_font
        cell22.alignment = side_alignment
        row = 2
        col = 2

        for slip in self.slip_ids:
            if slip.employee_id.bank_name_id != self.env.company.bank_id:
                worksheet_lbt.cell(row=row, column=col, value="LBT")
                col += 1
                worksheet_lbt.cell(row=row, column=col, value=account)
                col += 1

                worksheet_lbt.cell(row=row, column=col, value="")
                col += 1
                worksheet_lbt.cell(row=row, column=col, value=slip.employee_id.iban_num or "")
                col += 1
                worksheet_lbt.cell(row=row, column=col, value=slip.employee_id.display_name)
                col += 1
                worksheet_lbt.cell(row=row, column=col, value=slip.employee_id.addresline1 or "")
                col += 1
                worksheet_lbt.cell(row=row, column=col, value=slip.employee_id.addresline2 or "")
                col += 1

                worksheet_lbt.cell(row=row, column=col,
                                   value=slip.employee_id.bank_name_id.name or "")
                col += 1

                worksheet_lbt.cell(row=row, column=col, value="AE")
                col += 1
                worksheet_lbt.cell(row=row, column=col, value="UAE")
                col += 1
                worksheet_lbt.cell(row=row, column=col, value="UAE")
                col += 1

                worksheet_lbt.cell(row=row, column=col,
                                   value=slip.employee_id.swift_address_bic or "")
                col += 1
                worksheet_lbt.cell(row=row, column=col, value="")

                col += 1
                worksheet_lbt.cell(row=row, column=col, value=self.date_end.strftime('%d/%m/%Y'))
                col += 1
                worksheet_lbt.cell(row=row, column=col,
                                   value=self.env.company.currency_id.display_name)
                col += 1

                worksheet_lbt.cell(row=row, column=col, value=slip.net_wage)
                col += 1

                worksheet_lbt.cell(row=row, column=col, value="OUR")
                col += 1
                worksheet_lbt.cell(row=row, column=col, value="")
                col += 1

                worksheet_lbt.cell(row=row, column=col,
                                   value=self.date_end.strftime("%B %Y") + " Salary")
                col += 1

                col = 2
                row += 1
        # Save the workbook to a BytesIO object
        report_file = BytesIO()
        workbook.save(report_file)
        report_file.seek(0)

        filename = 'Non WPS ' + self.date_end.strftime("%B %Y") + ".xlsx"
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'datas': base64.encodebytes(report_file.read()),
            'res_model': self._name,
            'res_id': self.id
        })

        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s?download=true' % attachment.id,
            'target': 'new',
        }

    def print_xlsx_report_kuec(self):
        """
                Generates a WPS (Wage Protection System) Excel report listing employee payment details.
                Fields include employee unique ID, routing number, IBAN, pay period, salary, and unpaid days.
                """
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = 'WPS Report'

        # for col_num in range(1, 9):
        #     worksheet.col(col_num).width = 10000
        column_widths = [30, 30, 30, 30, 30, 30, 30, 30]
        for i, width in enumerate(column_widths, start=1):
            worksheet.column_dimensions[get_column_letter(i)].width = width

        headers = [
            "Employee Unique ID", "Agent ID/ Routing No.", "Employee Account with Agent Bank IBAN",
            "Pay Start Date", "Pay End Date", "Income Fixed Component", "Income Variable Component",
            "Days on Leave for Period"
        ]
        header_fill = PatternFill(start_color="ccdfe0", end_color="ccdfe0", fill_type="solid")
        header_font = Font(bold=True)
        header_style = xlwt.easyxf(
            'font: bold on; pattern: pattern solid, fore_colour ocean_blue; align: horiz center;')
        for col_num, header in enumerate(headers):
            cell = worksheet.cell(row=1, column=col_num + 1, value=header)
            # cell(row=1, column=1, value="1. Own Account Transfer \n \n 2.Transfer Within Bank")

            cell.fill = header_fill
            cell.font = header_font

        # Fill in the sample data based on the image
        data = [
            ["Employee MOL", "Agent Id /Routing No", "Employee Bank IBAN", "Pay Start Date",
             "Pay End Date", " Total Salary", "Overtime", "Days on Leave for Period"]
        ]

        # Add data to the worksheet
        # for row_num, row_data in enumerate(data, 1):
        #     for col_num, cell_value in enumerate(row_data):
        #         worksheet.write(row_num, col_num, cell_value)
        row = 2
        col = 1
        if self.slip_ids:
            for slip in self.slip_ids:

                num_ofdays = 0
                if slip.employee_id:
                    num_ofdays = self.get_employee_unpaid_timeoff(slip.employee_id.id)
                worksheet.cell(row, col, slip.employee_id.emp_unique_id or "")
                col += 1
                worksheet.cell(row, col, slip.employee_id.emp_routing_no or "")
                col += 1
                worksheet.cell(row, col, slip.employee_id.iban_num or "")
                col += 1
                worksheet.cell(row, col, self.date_start.strftime('%d/%m/%Y'))
                col += 1
                worksheet.cell(row, col, self.date_end.strftime('%d/%m/%Y'))
                col += 1
                worksheet.cell(row, col, slip.net_wage)

                col += 2
                worksheet.cell(row, col, num_ofdays)
                row += 1
                col = 1

        report_file = BytesIO()
        workbook.save(report_file)
        report_file.seek(0)

        filename = 'WPS ' + self.date_end.strftime("%B %Y") + ".xlsx"
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'datas': base64.encodebytes(report_file.read()),
            'res_model': self._name,
            'res_id': self.id
        })

        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s?download=true' % attachment.id,
            'target': 'new',
        }

    def print_registration_report_kuec(self):
        """
                Generates a detailed Payroll Register Excel report with:
                - Employee and contract metadata (KU #, job, department, grade, nationality).
                - Earnings and deductions (auto-mapped from salary rules).
                - Gross salary, deductions, and net salary summary.

                Salary rules under categories 'ALW', 'BASIC', and 'DED' are dynamically added to the header.
                """
        workbook = openpyxl.Workbook()
        ws = workbook.active
        ws.title = 'Payroll Register'

        # Set column widths
        for i in range(1, 51):
            ws.column_dimensions[get_column_letter(i)].width = 30

        # Style definitions
        side_font = Font(color="FFFFFF", bold=True)
        side_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header_style = {
            "font": Font(bold=True),
            "alignment": Alignment(horizontal="center", vertical="center"),
            "fill": PatternFill(start_color="0155b8", end_color="0155b8", fill_type="solid"),
            "border": Border(
                left=Side(style='thin', color='000000'),
                right=Side(style='thin', color='000000'),
                top=Side(style='thin', color='000000'),
                bottom=Side(style='thin', color='000000')
            )
        }

        def apply_header_style(cell):
            cell.font = header_style["font"]
            cell.alignment = header_style["alignment"]
            cell.fill = header_style["fill"]
            cell.border = header_style["border"]

        # Top Cell
        ws.merge_cells('A1:M1')
        cell = ws.cell(row=1, column=1,
                       value=f"Payroll Period:-{self.date_end.strftime('%d/%m/%Y')}")
        cell.font = side_font
        cell.alignment = side_alignment
        cell.border = header_style["border"]

        # Merge Header Columns
        for col in range(1, 11):
            ws.merge_cells(start_row=2, start_column=col, end_row=4, end_column=col)

        headers_base = [
            "KU Employee Number", "Ankabut ID", "Name", "Joining Date", "Position",
            "Section", "Grade", "Nationality", "Contract End Date", "Last Working Date"
        ]

        value_index_map = {}

        for col, title in enumerate(headers_base, start=1):
            cell = ws.cell(row=2, column=col, value=title)
            apply_header_style(cell)

        if not self.slip_ids:
            return workbook

        row = 5
        structure = self.slip_ids[0].struct_id

        # Earnings and Deductions
        def build_rule_section(code_list, start_col, title):
            rules = structure.rule_ids.filtered(
                lambda r: r.category_id and r.category_id.code in code_list)
            rule_names = rules.mapped("name")
            if rule_names:
                end_col = start_col + len(rule_names) - 1
            else:
                end_col = start_col
            ws.merge_cells(start_row=2, start_column=start_col, end_row=2, end_column=end_col)
            cell = ws.cell(row=2, column=start_col, value=title)
            apply_header_style(cell)

            for idx, name in enumerate(rule_names, start=start_col):
                ws.merge_cells(start_row=3, start_column=idx, end_row=4, end_column=idx)
                cell = ws.cell(row=3, column=idx, value=name)
                value_index_map[name] = idx
                apply_header_style(cell)
            return rule_names, end_col

        earnings_headers, earnings_end = build_rule_section(['ALW', 'BASIC'], 11, "Earnings")
        deductions_headers, deductions_end = build_rule_section(['DED'], earnings_end + 1,
                                                                "Voluntary Deductions")

        summary_start = deductions_end + 1
        summary_titles = ["Gross Salary", "Net Deduction", "Total"]
        for offset, title in enumerate(summary_titles):
            col = summary_start + offset
            ws.merge_cells(start_row=2, start_column=col, end_row=4, end_column=col)
            value_index_map[title] = col
            cell = ws.cell(row=2, column=col, value=title)
            apply_header_style(cell)

        # Data rows
        for slip in self.slip_ids:
            emp = slip.employee_id
            contract = slip.contract_id

            row_data = [
                emp.ku_number or "", emp.sequence or "", emp.display_name or "",
                contract.ku_hire_date.strftime(
                    '%d/%m/%Y') if contract and contract.ku_hire_date else "",
                contract.job_id.display_name if contract and contract.job_id else "",
                contract.department_id.display_name if contract and contract.department_id else "",
                contract.grade_id.display_name if contract and contract.grade_id else "",
                emp.country_id.display_name or "",
                contract.date_end.strftime('%d/%m/%Y') if contract and contract.date_end else "",
                ""  # Last Working Date
            ]

            for col, val in enumerate(row_data, start=1):
                ws.cell(row=row, column=col, value=val)

            gross_amount = 0
            net_deduction_amount = 0
            for line in slip.line_ids.filtered(
                    lambda x: x.category_id.code not in ['GROSS', 'NET']):
                if line.category_id.code in ['ALW', 'BASIC']:
                    gross_amount += line.total
                elif line.category_id.code in ['DED']:
                    net_deduction_amount += line.total
                ws.cell(row=row, column=value_index_map.get(line.name, summary_start + 4),
                        value=f"-{round(line.total, 2):,}" if line.category_id.code in [
                            'DED'] and line.total != 0 else f"{round(line.total, 2):,}")
            total_amount = gross_amount - net_deduction_amount
            ws.cell(row=row, column=value_index_map.get('Gross Salary'),
                    value=f"{round(gross_amount, 2):,}")
            ws.cell(row=row, column=value_index_map.get('Net Deduction'),
                    value=f"-{round(net_deduction_amount, 2):,}" if net_deduction_amount != 0 else f"{round(net_deduction_amount, 2):,}")
            ws.cell(row=row, column=value_index_map.get('Total'),
                    value=f"{round(total_amount, 2):,}")
            row += 1

        report_file = BytesIO()
        workbook.save(report_file)
        report_file.seek(0)

        filename = 'Payroll Register ' + self.date_end.strftime("%B %Y") + ".xlsx"
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'datas': base64.encodebytes(report_file.read()),
            'res_model': self._name,
            'res_id': self.id
        })

        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s?download=true' % attachment.id,
            'target': 'new',
        }





