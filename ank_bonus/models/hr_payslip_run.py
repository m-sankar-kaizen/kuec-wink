# -*- coding: utf-8 -*-
import openpyxl
import base64

from io import BytesIO
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter

from odoo import models


class HrPayslipRun(models.Model):
    """
    Inherits the hr.payslip.run model to provide functionality
    for generating a detailed Bonus Report in Excel format.

    The generated report includes:
    - Employee number, name, and personal details.
    - Job and contract-related information.
    - Bonus, gross salary, deductions, and net total.

    This report is saved as an attachment and available for download via URL.
    """
    _inherit = 'hr.payslip.run'

    def print_bonus(self):
        """
        Generate a Bonus Report in Excel format for the current payslip batch.
        The report contains key payroll details for each employee:
        - Personal & employment details
        - Bonus components
        - Salary and deduction breakdown

        Returns:
            dict: An action URL that triggers a download of the Excel file.
        """
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = 'Bonus Report'

        column_widths = [30, 30, 30, 30, 30, 30, 30, 30]
        for i, width in enumerate(column_widths, start=1):
            worksheet.column_dimensions[get_column_letter(i)].width = width

        headers = [
            "Employee Number", "Name", "Joining Date", "Position", "Section",
            "Grade",
            "Nationality", "Contract End Date", "Earnings/ Bonus and ex gratia",
            "Gross Salary", "Net Deduction", "Total"
        ]

        header_fill = PatternFill(start_color="ccdfe0", end_color="ccdfe0",
                                  fill_type="solid")
        header_font = Font(bold=True)

        for col_num, header in enumerate(headers):
            cell = worksheet.cell(row=1, column=col_num + 1, value=header)
            cell.fill = header_fill
            cell.font = header_font

        row = 2
        if self.slip_ids:
            for slip in self.slip_ids:
                col = 1
                worksheet.cell(row, col, slip.employee_id.emp_unique_id or "");
                col += 1
                worksheet.cell(row, col, slip.employee_id.name or "");
                col += 1
                worksheet.cell(row, col, slip.employee_id.hire_date.strftime(
                    '%d/%m/%Y') if slip.employee_id.hire_date else "");
                col += 1
                worksheet.cell(row, col, slip.employee_id.job_id.name or "");
                col += 1
                worksheet.cell(row, col,
                               slip.employee_id.section_id.name or "");
                col += 1
                worksheet.cell(row, col, slip.employee_id.grade_id.name or "");
                col += 1
                worksheet.cell(row, col,
                               slip.employee_id.country_id.name or "");
                col += 1
                worksheet.cell(row, col, slip.contract_id.date_end.strftime(
                    '%d/%m/%Y') if slip.contract_id.date_end else "");
                col += 1

                # Bonus
                bonus = sum(slip.line_ids.filtered(
                    lambda x: x.category_id == self.env.ref(
                        'ank_bonus.BONUS')).mapped('total'))
                worksheet.cell(row, col, bonus or "");
                col += 1

                # Gross Salary
                gross_salary = sum(slip.line_ids.filtered(
                    lambda x: x.category_id == self.env.ref(
                        'hr_payroll.GROSS')).mapped('total'))
                worksheet.cell(row, col, gross_salary or "");
                col += 1

                # Net Deduction
                net_ded = sum(slip.line_ids.filtered(
                    lambda x: x.category_id == self.env.ref(
                        'hr_payroll.DED')).mapped('total'))
                worksheet.cell(row, col, net_ded or "");
                col += 1

                # Net Total
                net_wage = sum(slip.line_ids.filtered(
                    lambda x: x.category_id == self.env.ref(
                        'hr_payroll.NET')).mapped('total'))
                worksheet.cell(row, col, net_wage or "")

                row += 1

        # Save and encode the report
        report_file = BytesIO()
        workbook.save(report_file)
        report_file.seek(0)

        filename = 'BONUS ' + self.date_end.strftime("%B %Y") + ".xlsx"
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'datas': base64.encodebytes(report_file.read()),
            'res_model': self._name,
            'res_id': self.id
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'new',
        }
