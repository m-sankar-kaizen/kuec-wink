# -*- coding: utf-8 -*-
import openpyxl
import base64

from io import BytesIO
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

from odoo import models


class HrPayslipRun(models.Model):
    """
        Inherits the hr.payslip.run model to add functionality for generating
        GCC and UAE pension Excel reports based on payroll data for each employee.

        These reports are tailored to the requirements of the Abu Dhabi Retirement Pensions
        and Benefits Fund and include relevant salary and pension contribution breakdowns.
        """
    _inherit = 'hr.payslip.run'

    def _get_gcc_pension_report(self):
        """
                Prepares the pension data for employees from GCC countries excluding UAE.

                :return: List of dictionaries containing report data for each employee.
                :rtype: list[dict]
                """
        gcc_countries = ['BH', 'KW', 'OM', 'QA', 'SA']  # No need to include UAE (AE)
        records = self.slip_ids.filtered(lambda x: x.employee_id.country_id.code in gcc_countries)
        datas = []
        for rec in records:
            basic_salary = sum(rec.line_ids.filtered(lambda x: x.code == 'BASIC').mapped('amount'))
            living_allowance = sum(rec.line_ids.filtered(lambda x: x.code == 'LIVALL').mapped('amount'))
            employee_share = sum(rec.line_ids.filtered(lambda x: x.code in ['PCOMAN', 'PEAOALL', 'PECOALL']).mapped('amount'))
            company_share = sum(rec.line_ids.filtered(lambda x: x.code == 'CCEO').mapped('amount'))
            data = {
                'ku_number': rec.employee_id.ku_number,
                'ank_number': rec.employee_id.sequence,
                'name': rec.employee_id.name,
                'grade': rec.employee_id.grade_id.name,
                'contract_type': rec.struct_id.name,
                'department': rec.employee_id.department_id.name,
                'cost_center': False,
                'basic_salary': round(basic_salary, 2),
                'living_allowance': round(living_allowance, 2),
                'employee_share': round(employee_share, 2),
                'company_share': round(company_share, 2),
                'total_share': round(company_share + employee_share, 2),
            }
            datas.append(data)

        return datas

    def print_gcc_pension_report(self):
        """
                Generates the GCC pension Excel report, creates an attachment, and returns a URL action to download it.

                :return: Dictionary with ir.actions.act_url to download the report.
                :rtype: dict
                """
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = 'Abu Dhabi Retirement Pensions and Benefits Fund - Salary Reporting'

        # Set column widths
        column_widths = [25, 25, 40, 20, 25, 35, 30, 30, 30, 30, 30, 30]
        for i, width in enumerate(column_widths, start=1):
            worksheet.column_dimensions[get_column_letter(i)].width = width

        # Headers
        headers = [
            "Employee Number", "Ank ID", "Employee Name", "Grade",
            "Contract Type", "Department", "Cost Center",
            "Basic Salary", "Living Allowance", "Employee Share",
            "Entity Share", "Total Share"]

        # Headers (row 2)
        header_fill = PatternFill(start_color="d3d3d3", end_color="d3d3d3", fill_type="solid")
        bold = Font(bold=True)
        center_align = Alignment(horizontal="center", vertical="center")
        right_center_align = Alignment(horizontal="right", vertical="center")

        for col_num, header in enumerate(headers, start=1):  # start from column 1 now
            cell = worksheet.cell(row=2, column=col_num, value=header)
            cell.fill = header_fill
            cell.font = bold
            cell.alignment = center_align

        report = self._get_gcc_pension_report()

        sum_employee_share = 0
        sum_company_share = 0
        sum_total_share = 0

        row = 3

        for rec in report:
            for col_index, key in enumerate(rec.keys(), start=1):
                cell = worksheet.cell(row=row, column=col_index, value=rec[key])
                if key in ['ku_number', 'ank_number', 'grade', 'contract_type', 'department']:
                    cell.alignment = center_align
                else:
                    cell.number_format = '#,##0.00'
            sum_employee_share += rec['employee_share']
            sum_company_share += rec['company_share']
            sum_total_share += rec['total_share']

            row += 1

        # FOOTER
        worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        cell = worksheet.cell(row=row, column=1, value="TOTAL")
        cell.font = bold
        cell.alignment = right_center_align
        cell.fill = header_fill

        worksheet.cell(row=row, column=10, value=round(sum_employee_share, 2))
        worksheet.cell(row=row, column=11, value=round(sum_company_share, 2))
        worksheet.cell(row=row, column=12, value=round(sum_total_share, 2))

        for col in range(10, 12):
            cell = worksheet.cell(row=row, column=col)
            cell.font = bold
            cell.number_format = '#,##0.00'

        report_file = BytesIO()
        workbook.save(report_file)
        report_file.seek(0)

        filename = 'GCC Pension Report ' + self.date_end.strftime("%B %Y") + ".xlsx"
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'datas': base64.encodebytes(report_file.read()),
            'res_model': self._name,
            'res_id': self.id
        })

        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s?download=true' % attachment.id,
            'target': 'new',
        }

    def _get_uae_pension_report(self):
        """
            Collect and compute the detailed UAE pension report data for all employees
            marked as 'local' in the payslip batch (i.e., UAE nationals).

            This method extracts the necessary salary components from the payslip lines
            such as basic salary, allowances, pension shares, and computes:
            - Total pensionable salary
            - Total shares (employee + employer)
            - Gross and net salary

            Returns:
                list[dict]: A list of dictionaries, where each dictionary contains
                            data for one employee's pension report line.
            """
        records = self.slip_ids.filtered(lambda x: x.employee_id.kaz_employee_type == 'local')
        datas = []
        for rec in records:
            basic_salary = sum(rec.line_ids.filtered(lambda x: x.code == 'BASIC').mapped('amount'))
            child_allowance = sum(rec.line_ids.filtered(lambda x: x.code in ['CHLDALW', 'CHLALW']).mapped('amount'))
            premium_allowance = sum(rec.line_ids.filtered(lambda x: x.code == 'PREALW').mapped('amount'))
            living_allowance = sum(rec.line_ids.filtered(lambda x: x.code == 'LIVALL').mapped('amount'))
            personal_allowance = sum(rec.line_ids.filtered(lambda x: x.code == 'PERALW').mapped('amount'))
            connectivity_allowance = sum(rec.line_ids.filtered(lambda x: x.code == 'CONALW').mapped('amount'))
            employee_share = sum(rec.line_ids.filtered(lambda x: x.code in ['PEND', 'PECUALL']).mapped('amount'))
            company_share = sum(rec.line_ids.filtered(lambda x: x.code == 'CCONU').mapped('amount'))
            total_gross_salary = sum(rec.line_ids.filtered(lambda x: x.code == 'GROSS').mapped('amount'))
            total_net_salary = sum(rec.line_ids.filtered(lambda x: x.code == 'NET').mapped('amount'))
            total_pension_amount = basic_salary + child_allowance + premium_allowance + living_allowance
            data = {
                'emirates_id': rec.employee_id.emirates,
                'ku_number': rec.employee_id.ku_number,
                'ank_number': rec.employee_id.sequence,
                'name': rec.employee_id.name,
                'basic_salary': round(basic_salary, 2),
                'child_allowance': round(child_allowance, 2),
                'premium_allowance': round(premium_allowance, 2),
                'living_allowance': round(living_allowance, 2),
                'personal_allowance': round(personal_allowance, 2),
                'connectivity_allowance': round(connectivity_allowance, 2),
                'total_pension_amount': round(total_pension_amount, 2),
                'pension_percentage': 'PLAN5',
                'employee_share': round(employee_share, 2),
                'company_share': round(company_share, 2),
                'total_share': round(company_share + employee_share, 2),
                'total_gross_salary': round(total_gross_salary, 2),
                'net_payroll': round(total_net_salary, 2),
            }
            datas.append(data)

        return datas

    def print_uae_pension_report(self):
        """
           Generate a UAE Pension Report Excel file for all local employees (UAE nationals)
           based on payslip data in the current payslip run.

           The report includes:
           - Detailed breakdown of salary components per employee
           - Employer/Employee shares for pension
           - Totals for all monetary columns at the bottom
           - Saved as an attachment and triggers file download

           Returns:
               dict: An `ir.actions.act_url` dict that opens the generated Excel report in a new tab.
           """
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = 'Abu Dhabi Retirement Pensions and Benefits Fund - Salary Reporting'

        # Set column widths
        column_widths = [10, 25, 15, 15, 45, 25, 25, 40, 25, 25, 30, 30, 30, 30, 30, 30, 30, 30]
        for i, width in enumerate(column_widths, start=1):
            worksheet.column_dimensions[get_column_letter(i)].width = width

        # Headers
        headers = [
            "Sl No.", "National ID",
            "Emp No.", "Ank ID", "Employee Name",
            "Basic Salary", "Child Allowance",
            "Social Allowance and UAE Premium", "Living Allowance", "Personal Allowance",
            "Connectivity Allowance", "Total Pensionable Salary", "Pension Percentage",
            "Employee Share", "Company Share", "Total Contribution",
            "Total Gross Salary", "Net Payroll Amount"]

        side_fill = PatternFill(start_color="800000", end_color="800000", fill_type="solid")
        side_font = Font(color="FFFFFF", bold=True)
        side_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        worksheet.merge_cells('A1:D1')
        cell1 = worksheet.cell(row=1, column=1,
                               value="Abu Dhabi Retirement Pensions and Benefits Fund - Salary Reporting")
        cell1.alignment = side_alignment
        cell1.fill = side_fill
        cell1.font = side_font

        # Headers (row 2)
        header_fill = PatternFill(start_color="ccdfe0", end_color="ccdfe0", fill_type="solid")
        bold = Font(bold=True)
        center_align = Alignment(horizontal="center", vertical="center")

        for col_num, header in enumerate(headers, start=1):  # start from column 1 now
            cell = worksheet.cell(row=2, column=col_num, value=header)
            cell.fill = header_fill
            cell.font = bold
            cell.alignment = center_align

        report = self._get_uae_pension_report()

        sum_basic_salary = 0
        sum_child_allowance = 0
        sum_premium_allowance = 0
        sum_living_allowance = 0
        sum_personal_allowance = 0
        sum_connectivity_allowance = 0
        sum_total_pension_amount = 0
        sum_employee_share = 0
        sum_company_share = 0
        sum_total_share = 0
        sum_total_gross_salary = 0
        sum_net_payroll = 0

        row = 3

        for rec in report:
            worksheet.cell(row, 1, row - 2)
            for col_index, key in enumerate(rec.keys(), start=2):
                cell = worksheet.cell(row=row, column=col_index, value=rec[key])
                if key in ['ku_number', 'ank_number', 'pension_percentage']:
                    cell.alignment = center_align
                if key not in ['emirates_id', 'ku_number', 'ank_number', 'name', 'pension_percentage']:
                    cell.number_format = '#,##0.00'
            sum_basic_salary += rec['basic_salary']
            sum_child_allowance += rec['child_allowance']
            sum_premium_allowance += rec['premium_allowance']
            sum_living_allowance += rec['living_allowance']
            sum_personal_allowance += rec['personal_allowance']
            sum_connectivity_allowance += rec['connectivity_allowance']
            sum_total_pension_amount += rec['total_pension_amount']
            sum_employee_share += rec['employee_share']
            sum_company_share += rec['company_share']
            sum_total_share += rec['total_share']
            sum_total_gross_salary += rec['total_gross_salary']
            sum_net_payroll += rec['net_payroll']

            row += 1

        # FOOTER
        cell = worksheet.cell(row=row, column=5, value="TOTAL")
        cell.font = bold
        cell.alignment = center_align
        worksheet.cell(row=row, column=6, value=round(sum_basic_salary, 2))
        worksheet.cell(row=row, column=7, value=round(sum_child_allowance, 2))
        worksheet.cell(row=row, column=8, value=round(sum_premium_allowance, 2))
        worksheet.cell(row=row, column=9, value=round(sum_living_allowance, 2))
        worksheet.cell(row=row, column=10, value=round(sum_personal_allowance, 2))
        worksheet.cell(row=row, column=11, value=round(sum_connectivity_allowance, 2))
        worksheet.cell(row=row, column=12, value=round(sum_total_pension_amount, 2))
        # Column 13 is pension percentage – skip
        worksheet.cell(row=row, column=14, value=round(sum_employee_share, 2))
        worksheet.cell(row=row, column=15, value=round(sum_company_share, 2))
        worksheet.cell(row=row, column=16, value=round(sum_total_share, 2))
        worksheet.cell(row=row, column=17, value=round(sum_total_gross_salary, 2))
        worksheet.cell(row=row, column=18, value=round(sum_net_payroll, 2))

        for col in range(6, 19):
            cell = worksheet.cell(row=row, column=col)
            cell.font = bold
            cell.number_format = '#,##0.00'

        report_file = BytesIO()
        workbook.save(report_file)
        report_file.seek(0)

        filename = 'UAE Pension Report ' + self.date_end.strftime("%B %Y") + ".xlsx"
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'datas': base64.encodebytes(report_file.read()),
            'res_model': self._name,
            'res_id': self.id
        })

        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s?download=true' % attachment.id,
            'target': 'new',
        }
